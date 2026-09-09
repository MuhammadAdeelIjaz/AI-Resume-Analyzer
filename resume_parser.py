from io import BytesIO


class ResumeParserError(Exception):
    """Custom exception for resume parsing errors."""


class ResumeParser:
    """
    Responsible only for extracting text from
    supported resume file formats.
    """

    # =====================================================
    # Public Method
    # =====================================================

    @staticmethod
    def extract_text(
        file_bytes: bytes,
        filename: str,
    ) -> str:

        if not file_bytes:

            raise ResumeParserError(
                "The uploaded file is empty."
            )

        extension = (
            filename
            .lower()
            .rsplit(".", 1)[-1]
            if "." in filename
            else ""
        )

        if extension == "pdf":

            return ResumeParser._extract_pdf(
                file_bytes
            )

        if extension == "docx":

            return ResumeParser._extract_docx(
                file_bytes
            )

        raise ResumeParserError(
            "Unsupported file type. "
            "Please upload a PDF or DOCX file."
        )

    # =====================================================
    # PDF Extraction
    # =====================================================

    @staticmethod
    def _extract_pdf(
        file_bytes: bytes,
    ) -> str:

        try:

            import pypdf

            reader = pypdf.PdfReader(
                BytesIO(file_bytes)
            )

            pages = []

            for page in reader.pages:

                text = (
                    page.extract_text()
                    or ""
                )

                if text.strip():

                    pages.append(
                        text.strip()
                    )

            result = "\n\n".join(
                pages
            ).strip()

            if not result:

                raise ResumeParserError(
                    "No readable text was found "
                    "in the PDF."
                )

            return result

        except ResumeParserError:

            raise

        except Exception as error:

            raise ResumeParserError(
                f"Could not read PDF: {error}"
            )

    # =====================================================
    # DOCX Extraction
    # =====================================================

    @staticmethod
    def _extract_docx(
        file_bytes: bytes,
    ) -> str:

        try:

            from docx import Document

            document = Document(
                BytesIO(file_bytes)
            )

            parts = []

            # ---------------------------------------------
            # Paragraphs
            # ---------------------------------------------

            for paragraph in document.paragraphs:

                text = paragraph.text.strip()

                if text:

                    parts.append(text)

            # ---------------------------------------------
            # Tables
            # ---------------------------------------------

            for table in document.tables:

                for row in table.rows:

                    cells = []

                    for cell in row.cells:

                        text = (
                            cell.text
                            .strip()
                        )

                        cells.append(text)

                    if any(cells):

                        parts.append(
                            " | ".join(cells)
                        )

            result = "\n".join(
                parts
            ).strip()

            if not result:

                raise ResumeParserError(
                    "No readable text was found "
                    "in the DOCX file."
                )

            return result

        except ResumeParserError:

            raise

        except Exception as error:

            raise ResumeParserError(
                f"Could not read DOCX: {error}"
            )